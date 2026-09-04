import io
from pathlib import Path

import pytest
from conftest import make_pdf
from docx import Document

from agent import AgentConfig, create_demo_agent
from agent.tools.documents import parse_document


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "gb18030", "utf-16"])
def test_txt_encodings_and_safe_filename(demo_agent, cardiac_text, encoding):
    result = demo_agent.classify_file(cardiac_text.encode(encoding), "../../论文.TXT")
    assert result.status == "success"
    assert result.prediction.label_id == 0
    assert result.input.filename == "论文.TXT"
    assert "parse_document" in [step.step for step in result.steps]
    assert "read_text" not in [step.step for step in result.steps]


def test_docx_paragraph_table_order_and_end_to_end(demo_agent):
    document = Document()
    document.add_paragraph("标题：心血管研究")
    document.add_paragraph("摘要：高血压患者心血管风险与冠心病有关。")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "关键词：高血压；冠心病"
    document.add_paragraph("参考文献")
    stream = io.BytesIO()
    document.save(stream)
    parsed = parse_document(stream.getvalue(), "paper.docx", AgentConfig())
    assert parsed.text.index("摘要") < parsed.text.index("关键词") < parsed.text.index("参考文献")
    result = demo_agent.classify_file(stream.getvalue(), "paper.docx")
    assert result.status == "success"
    assert result.input.preparation_strategy == "document_fields"
    assert result.prediction.label_id == 0


def test_pdf_with_text_and_partial_blank_pages(demo_agent):
    result = demo_agent.classify_file(
        make_pdf("A study of hypertension and cardiovascular disease.", blank_pages=1), "paper.pdf"
    )
    assert result.status == "success"
    assert result.input.pages == 2
    assert result.prediction.label_id == 0
    assert "PARTIAL_PDF_TEXT" in [note.code for note in result.warnings]


@pytest.mark.parametrize(
    "content,filename,code",
    [
        (b"", "empty.txt", "EMPTY_FILE"),
        (b"plain text", "wrong.pdf", "INVALID_FILE"),
        (b"not a zip", "wrong.docx", "DOCUMENT_PARSE_FAILED"),
        (b"text text", "wrong.exe", "UNSUPPORTED_FILE_TYPE"),
        (make_pdf(), "scanned.pdf", "NO_EXTRACTABLE_TEXT"),
        (make_pdf("Private text", encrypted=True), "encrypted.pdf", "ENCRYPTED_PDF"),
    ],
)
def test_file_failures_are_structured(demo_agent, content, filename, code):
    result = demo_agent.classify_file(content, filename)
    assert result.status == "error"
    assert result.error.code == code
    assert "classify_text" not in [step.step for step in result.steps]


def test_file_size_and_page_limits():
    agent = create_demo_agent(config=AgentConfig(max_file_bytes=4))
    assert agent.classify_file(b"12345", "x.txt").error.code == "FILE_TOO_LARGE"
    agent = create_demo_agent(config=AgentConfig(max_pdf_pages=1))
    assert agent.classify_file(make_pdf(blank_pages=1), "x.pdf").error.code == "TOO_MANY_PAGES"


def test_extracted_document_fields_match_structured_input(demo_agent):
    root = Path(__file__).resolve().parents[1]
    import json

    fields = json.loads((root / "examples/sample_paper.json").read_text())
    uploaded = demo_agent.classify_file((root / "examples/sample.txt").read_bytes(), "sample.txt")
    direct = demo_agent.classify_paper(**fields)
    assert uploaded.input.preparation_strategy == "document_fields"
    assert [item.score for item in uploaded.candidates] == [
        item.score for item in direct.candidates
    ]
